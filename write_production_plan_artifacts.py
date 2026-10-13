"""Create editable rollout plan + practical handover templates; no runtime edits."""
import csv
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'production_plan'
OUT.mkdir(exist_ok=True)

def csvfile(name,headers,rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(rows)

csvfile('BACKLOG_TRIEN_KHAI.csv',
 ['ID','Muc_uu_tien','Cong_viec','Nguoi_phu_trach_de_xuat','Phu_thuoc','Dieu_kien_hoan_thanh','Trang_thai'],[
 ['P0-01','P0','Xác định VPS từ chối kết nối, kiểm VM/firewall/gia hạn và backup','CNTT + chủ hosting','','Biên bản nguyên nhân; backup/snapshot trước khôi phục','Chưa làm'],
 ['P0-02','P0','Chốt use case, chủ dữ liệu và ma trận quyền pilot','Người bảo trợ + đầu mối TS/ĐT','','6–8 use case, 5–10 người, scope/route/owner rõ','Chờ trường xác nhận'],
 ['P0-03','P0','Private Git, loại secrets khỏi source, pin deps, staging và rollback','Nhóm + CNTT','P0-01','Build tái lập và test nền đạt trên Linux','Chưa làm'],
 ['P0-04','P0','SSO/OIDC, principal, quyền theo đơn vị/phạm vi/hành động','Backend + CNTT','P0-02,P0-03','Test sai quyền/tài khoản/thu hồi đạt','Chưa làm'],
 ['P0-05','P0','POST Run + UUID/owner, bỏ replay fallback production','Backend','P0-03,P0-04','Không câu hỏi trong URL; lịch sử khác người cách ly','Chưa làm'],
 ['P0-06','P0','Security envelope kế thừa xuống AI con/công cụ','Backend','P0-04','Constraints không bị bỏ; ma_gv lấy từ quyền server','Chưa làm'],
 ['P0-07','P0','Cổng kiểm payload/API routes trước planner và mọi model','Backend + phụ trách dữ liệu','P0-02,P0-06','Nguồn chưa phân loại cách ly; local-only không ra cloud','Chưa làm'],
 ['D-01','P0','Hợp đồng dataset và dictionary KPI TS/ĐT/liên phòng','Đầu mối dữ liệu + data engineer','P0-02','Chủ dữ liệu duyệt schema, mốc, định nghĩa, quyền','Chờ trường xác nhận'],
 ['D-02','P1','Snapshot/adapter chỉ đọc, staging đối soát và rollback','Data/backend','D-01,P0-07','KPI khớp báo cáo chính thức cùng snapshot','Chưa làm'],
 ['D-03','P1','Bỏ hardcode năm/kỳ/khóa/40:60/70; version hóa quy tắc','Đầu mối ĐT + backend','D-01','Dùng dữ liệu kỳ mới, đúng quy chế từng khóa','Chưa làm'],
 ['D-04','P1','Nguồn có owner/scope/version/hiệu lực/trang-mục; publish sau duyệt','Data/backend','D-01','Nguồn chưa duyệt không được truy xuất; recall giữ sửa đổi','Chưa làm'],
 ['R-01','P0','Postgres Run/state/events và checkpointer + queue worker','Backend','P0-03,P0-05','Kill/restart, hai worker trùng, retry không publish trùng','Chưa làm'],
 ['R-02','P1','Hủy/resume và SSE reconnect kiểm quyền','Backend + frontend','R-01,P0-06','Ngắt trình duyệt không mất Run; cancel chặn bước chưa gọi','Chưa làm'],
 ['Q-01','P1','Hợp đồng kết quả completeness/assumptions/evidence','Backend + nghiệp vụ','D-02,P0-06','Không cắt JSON; giữ cảnh báo; query lỗi báo partial','Chưa làm'],
 ['Q-02','P1','Final evidence và phê duyệt đầu ra ảnh hưởng con người','Backend + nghiệp vụ','Q-01,D-04','Sai mã/số/kỳ không gắn verified; có người duyệt','Chưa làm'],
 ['Q-03','P1','Server export phân trang, đủ dòng, chống công thức CSV','Backend + frontend','P0-04,D-02','Xuất >1000 dòng đúng quyền/số lượng, có audit','Chưa làm'],
 ['L-01','P1','Latency/cost từng attempt, deadline, connection pooling','Backend','P0-03','Ghi cả lỗi/retry; số liệu cost chưa biết hiển thị đúng','Chưa làm'],
 ['L-02','P1','Router hỏi lại/ngoài phạm vi và đường nhanh có evidence','Backend','Q-01,L-01','Không chạy fullgraph cho lời chào; case mơ hồ hỏi lại','Chưa làm'],
 ['L-03','P1','Chia đoạn/giảm payload, bỏ tổng hợp hoặc review trùng có điều kiện','Backend + QA','D-04,Q-02,L-01','A/B chất lượng không giảm; đo Run wallclock','Chưa làm'],
 ['E-01','P0','Bộ nghiệm thu mới; sửa grader và mẫu số tính lỗi','QA + người thẩm định','P0-02,D-01','Đáp án sai có số đúng bị trượt; lỗi gateway tính reliability','Chưa làm'],
 ['O-01','P0','Tài khoản dịch vụ, secrets, giám sát, backup/restore','CNTT + backend','P0-03,R-01','Restore thử; cảnh báo/owner/rollback rõ','Chưa làm'],
 ['M-01','P1','Gói NCKH: nguồn bài/case/DOI/định mức và worker','KH + tài chính + nhóm','D-01,P0-07','Kịch bản 3 case có bằng chứng; thiếu thì báo thiếu','Chờ dữ liệu'],
 ['M-02','P1','Gói HTQT: đối tác/thỏa thuận/chương trình/hiệu lực/quyền','HTQT + nhóm','D-01,P0-07','Truy vấn có nguồn/hiệu lực và bộ test riêng','Chờ dữ liệu'],
 ['PILOT','P0','Shadow rồi pilot có người duyệt 5–10 cán bộ','Nhà trường + nhóm','P0-04,P0-07,R-02,Q-02,Q-03,E-01,O-01','Qua cổng A–F trong kế hoạch; không lỗi nghiêm trọng mở','Chưa làm'],
 ['EXPAND','P2','Hành chính/Nhân sự/self-service và báo cáo sáu lĩnh vực','Chủ lĩnh vực + CNTT + nhóm','PILOT,M-01,M-02','Từng nghiệp vụ có dữ liệu/quyền/test/chuyển giao riêng','Chưa làm'],
])

csvfile('MAU_HOP_DONG_DU_LIEU.csv',
 ['dataset_id','owner_unit','steward','approved_by','schema_version','source_system','record_key','snapshot_id','snapshot_at','effective_from','effective_to','classification','allowed_roles','allowed_actions','allowed_fields','allowed_model_routes','retention_policy','quality_checks','checksum','status'],
 [[key,owner,'','','','','','','','','','','','','','','','','','Chờ duyệt'] for key,owner in [('TS-SNAPSHOT','Đơn vị Tuyển sinh'),('DT-SNAPSHOT','Đơn vị Đào tạo'),('NC-SOURCES','Đơn vị Khoa học'),('HTQT-SOURCES','Đơn vị HTQT')]])

csvfile('MAU_MA_TRAN_QUYEN.csv',
 ['role','resource','scope','action','allowed_fields','external_processing_allowed','approval_owner','status'],[
 ['Chuyên viên Tuyển sinh','Snapshot Tuyển sinh','Phạm vi được giao','read/aggregate','','','Chủ dữ liệu Tuyển sinh','Đề xuất - chưa cấp quyền'],
 ['Chuyên viên Tuyển sinh','Danh sách liên hệ','Phạm vi được giao','export','','','Chủ dữ liệu Tuyển sinh','Đề xuất - chưa cấp quyền'],
 ['Giảng viên','Lớp học phần','Lớp có phân công giảng dạy','read/aggregate','','','Chủ dữ liệu Đào tạo','Đề xuất - chưa cấp quyền'],
 ['Đại diện lãnh đạo','Báo cáo liên phòng','Tập nguồn được cấp riêng','aggregate/read','','','Chủ dữ liệu từng lĩnh vực','Đề xuất - chưa cấp quyền'],
 ['Quản trị kỹ thuật','Cấu hình model/hạ tầng','Chỉ cấu hình được giao','administer','','','CNTT','Không mặc định được đọc hồ sơ nghiệp vụ'],
])

doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
for name in ['Title','Heading 1','Heading 2','Heading 3']:
    doc.styles[name].font.name='Calibri';doc.styles[name].font.color.rgb=RGBColor.from_string('245BBA')
def clean(t):
    t=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m:m[1]+' ('+m[2]+')',t)
    return t.replace('**','').replace('`','')
lines=(ROOT/'KE_HOACH_VAN_HANH_THAT.md').read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    l=lines[i]
    if not l.strip():i+=1;continue
    if l.startswith('```'):
        lang=l[3:];block=[];i+=1
        while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
        if lang=='mermaid':
            doc.add_paragraph('Người dùng + SSO → FastAPI có quyền → PostgreSQL/queue → LangGraph worker → nguồn/công cụ có phạm vi → cổng dữ liệu → API được duyệt hoặc model nội bộ → kết quả theo quyền.')
            doc.add_paragraph('Không có tuyến hợp lệ: dừng/chờ. Run/events/artifacts lưu bền và được kiểm quyền trước hiển thị/xuất.')
        else:doc.add_paragraph('\n'.join(block))
    elif l.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[clean(v.strip()) for v in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',c) for c in cells):rows.append(cells)
            i+=1
        if rows:
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
            for row in rows:
                for cell,value in zip(table.add_row().cells,row):cell.text=value
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:run.font.size=Pt(9)
        continue
    elif l.startswith('# '):doc.add_heading(clean(l[2:]),0)
    elif l.startswith('## '):doc.add_heading(clean(l[3:]),1)
    elif l.startswith('### '):doc.add_heading(clean(l[4:]),2)
    elif l.startswith('- '):doc.add_paragraph(clean(l[2:]),style='List Bullet')
    elif re.match(r'^\d+\. ',l):doc.add_paragraph(clean(re.sub(r'^\d+\. ','',l)),style='List Number')
    else:doc.add_paragraph(clean(l))
    i+=1
doc.core_properties.title='Credere — Kế hoạch đưa vào vận hành thật tại nhà trường'
doc.core_properties.author='Nhóm dự án Credere'
doc.save(ROOT/'KE_HOACH_VAN_HANH_THAT.docx')
print('Created editable plan and 3 CSV handover/backlog templates.')
