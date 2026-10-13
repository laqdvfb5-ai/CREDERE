"""Build the editable Credere presentation and the matching speaker script.

Run with ~/.credere/venv/Scripts/python.exe presentation/build_deck.py.
Evidence: CONTENT_FACTS.json; screenshots: shots/current/screenshots.json.
No model calls or network access are needed to build this deck.
"""
import json
import shutil
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE / 'shots' / 'current'
FACTS = json.loads((HERE / 'CONTENT_FACTS.json').read_text(encoding='utf-8'))
OUT = HERE / 'Credere_Trinh_bay_BGH.pptx'
FONT = 'Segoe UI'
NAVY, INK, MUTED = '142238', '23364F', '687A91'
BLUE, TEAL, AMBER, PURPLE = '245BBA', '40836F', 'A76B27', '7755A2'
PAPER, WHITE, LINE, SOFT = 'F6F7F9', 'FFFFFF', 'DFE5ED', 'EBF1F9'
LIGHT, PALE_GREEN, PALE_AMBER = 'AFBFDA', 'ECF5F0', 'FBF3E6'
DATE = '29.09.2026'
W, H = 13.333333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Credere — Trợ lý AI phối hợp cho Trường Đại học Nghệ An'
prs.core_properties.subject = 'Bản demo hiện tại, cơ chế giao việc và đề xuất thử nghiệm có kiểm soát'
prs.core_properties.author = 'Nhóm đề tài Credere'
prs.core_properties.keywords = 'Credere, NAU, tuyển sinh, đào tạo, đa tác tử'
INDEX = []
TEXT_BOXES = []


def rgb(value):
    return RGBColor.from_string(value)


def rect(s, x, y, w, h, fill=WHITE, line=None, radius=.12):
    obj = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    if line:
        obj.line.color.rgb = rgb(line)
        obj.line.width = Pt(.8)
    else:
        obj.line.fill.background()
    if radius:
        obj.adjustments[0] = radius
    obj.shadow.inherit = False
    return obj


def text(s, x, y, w, h, content, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=.025):
    obj = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    lines = content.split('\n') if isinstance(content, str) else content
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.1
        p.space_after = Pt(5 if len(lines) > 1 else 0)
        r = p.add_run()
        r.text = str(line)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    TEXT_BOXES.append({'slide': len(prs.slides), 'text': '\n'.join(lines), 'font': size,
                       'x': x, 'y': y, 'w': w, 'h': h})
    return obj


def line(s, x1, y1, x2, y2, color=LINE, width=1.4, arrow=False):
    obj = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.line.color.rgb, obj.line.width = rgb(color), Pt(width)
    obj.shadow.inherit = False
    if arrow:
        end = etree.SubElement(obj.line._get_or_add_ln(), qn('a:tailEnd'))
        end.set('type', 'triangle')
    return obj


def pill(s, x, y, w, label, fill=SOFT, color=BLUE, size=12):
    rect(s, x, y, w, .34, fill=fill, radius=.2)
    text(s, x+.09, y+.025, w-.18, .25, label, size, color, bold=True)


def circle(s, x, y, d, label, color=BLUE, fill=SOFT, size=18):
    obj = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    obj.fill.solid(); obj.fill.fore_color.rgb = rgb(fill); obj.line.fill.background()
    obj.shadow.inherit = False
    text(s, x, y, d, d, label, size, color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def brand(s, x=.6, y=.22, color=NAVY):
    text(s, x, y, 2.2, .38, 'credere.', 20, color, bold=True)


def slide(title, section, seconds, notes, sources=(), dark=False, backup=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY if dark else PAPER)
    fg, muted = (WHITE, LIGHT) if dark else (INK, MUTED)
    brand(s, color=WHITE if dark else NAVY)
    text(s, 7.8, .29, 4.9, .23, f'TRƯỜNG ĐẠI HỌC NGHỆ AN  /  {DATE}', 10, muted, align=PP_ALIGN.RIGHT)
    line(s, .62, .82, 12.71, .82, '34465F' if dark else LINE, .7)
    text(s, .64, 1.03, 12, .28, section.upper(), 11, '98B5E7' if dark else BLUE, bold=True)
    if title:
        text(s, .62, 1.43, 12.1, .73, title, 31, fg, bold=True)
    num = len(prs.slides)
    text(s, .64, 7.02, 10.9, .24, 'Credere · Bản thử nghiệm nghiệp vụ' + (' · PHỤ LỤC DỰ PHÒNG' if backup else ''), 10, muted)
    text(s, 12.05, 6.98, .65, .3, f'{num:02}', 12, muted, align=PP_ALIGN.RIGHT)
    note = notes + ('\n\nCăn cứ (dành cho người trình bày):\n' + '\n'.join(str(p) for p in sources) if sources else '')
    s.notes_slide.notes_text_frame.text = note
    INDEX.append({'slide':num,'title':title or 'Credere — trợ lý AI phối hợp','section':section,'seconds':seconds,
                  'notes':notes,'sources':list(sources),'backup':backup})
    return s


def metric(s, x, y, w, number, label, color=BLUE, fill=WHITE):
    rect(s, x, y, w, 1.20, fill, LINE)
    text(s, x+.2, y+.13, w-.4, .55, str(number), 28, color, bold=True)
    text(s, x+.2, y+.77, w-.4, .32, label, 12, MUTED)


def feature(s, x, y, w, number, title, body, color=BLUE, fill=WHITE, h=1.3):
    rect(s, x, y, w, h, fill, LINE)
    circle(s, x+.18, y+.19, .42, number, color, SOFT if color==BLUE else PALE_GREEN, 13)
    text(s, x+.78, y+.15, w-.98, .45, title, 18, INK, bold=True)
    text(s, x+.78, y+.68, w-.98, h-.78, body, 14, MUTED)


def screen(s, filename, x, y, w, h, caption=None):
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(f'Missing actual screenshot: {path}')
    im = Image.open(path)
    iw, ih = im.size
    scale = min(w/iw, h/ih)
    pw, ph = iw*scale, ih*scale
    rect(s, x+(w-pw)/2-.035, y+(h-ph)/2-.035, pw+.07, ph+.07, WHITE, LINE, radius=.03)
    s.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))
    if caption:
        text(s, x, y+h+.10, w, .45, caption, 10, MUTED)


def node(s, x, y, w, h, title, body, color=BLUE, dark=False):
    fill, fg, secondary = ('223750', WHITE, LIGHT) if dark else (WHITE, INK, MUTED)
    rect(s, x, y, w, h, fill, '4B6380' if dark else LINE)
    rect(s, x+.02, y+.13, .045, h-.26, color, radius=False)
    text(s, x+.18, y+.12, w-.36, .42, title, 19, fg, bold=True)
    text(s, x+.18, y+.64, w-.36, h-.72, body, 14, secondary)


example = FACTS['principal_example']
question = example['question']
trace_source = example['source']
app_facts = {a['key']: a for a in FACTS['runtime']['applications']}
ts, dt = app_facts['tuyen_sinh'], app_facts['dao_tao']

# 01 — concrete promise, without claiming production readiness.
s = slide(None, 'Bản demo cập nhật', 30,
    'Kính thưa quý thầy cô, nhóm em xin trình bày Credere: bản thử trợ lý AI cho tuyển sinh, đào tạo và công việc tổng hợp của lãnh đạo. '
    'Điểm nhóm muốn chứng minh hôm nay là một yêu cầu có thể được giao cho các AI lĩnh vực, kiểm tra rồi ghép thành kết quả có căn cứ. '
    'Văn bản trong kho là tài liệu công khai; số liệu nội bộ dùng trong demo là dữ liệu giả lập.',
    ['backend/app/apps.py:117', 'CONTENT_FACTS.json'], dark=True)
text(s, .72, 1.76, 7.3, 2.6, 'Một câu hỏi.\nNhiều AI cùng làm.', 45, WHITE, bold=True)
text(s, .76, 4.15, 6.5, 1.12, 'Trợ lý nghiệp vụ cho tuyển sinh,\nđào tạo và điều hành nhà trường.', 24, LIGHT)
pill(s, .8, 5.91, 4.45, 'BẢN THỬ · VĂN BẢN CÔNG KHAI + SỐ LIỆU GIẢ LẬP', '283D57', 'CFDDF2', 10)
node(s, 8.5, 1.88, 3.8, 1.24, 'AI Tổng hợp', 'Giao việc · đối chiếu · tổng hợp', '8BADE6', True)
line(s, 8.5, 2.53, 8.14, 2.53, '8BADE6', 1.6)
line(s, 8.14, 2.53, 8.14, 5.63, '8BADE6', 1.6)
line(s, 8.14, 4.2, 8.5, 4.2, '8BADE6', 1.6, True)
line(s, 8.14, 5.63, 8.5, 5.63, '8BADE6', 1.6, True)
node(s, 8.5, 3.63, 3.8, 1.15, 'AI Tuyển sinh', 'Văn bản và số liệu tuyển sinh', '75B79D', True)
node(s, 8.5, 5.06, 3.8, 1.15, 'AI Đào tạo', 'Quy chế và số liệu học vụ', 'DBB074', True)

# 02 — a real recorded request is the thread through the whole story.
s = slide('Một câu hỏi cần nhiều đầu mối', 'Bài toán xuất phát', 45,
    'Ví dụ, Hiệu trưởng muốn biết ngành nào đang thiếu chỉ tiêu và ở những ngành đó sinh viên các khóa trước học tập ra sao. '
    'Một phần dữ liệu nằm ở Tuyển sinh, phần khác ở Đào tạo. Muốn ghép được phải thống nhất ngành, khóa, thời điểm và cách tính. '
    'Đây là tình huống đã có trong hồ sơ chạy của bản thử, không phải số liệu tình hình thật của Trường.', [trace_source])
rect(s, .7, 2.43, 11.93, 2.0, WHITE, LINE)
text(s, 1.0, 2.7, 11.3, 1.45, '“Ngành nào thiếu chỉ tiêu tuyển sinh nhất?\nỞ những ngành đó, cảnh báo học vụ và rời trường ra sao?”', 27, INK, bold=True)
for x, n, h, b in [( .7,'01','Tuyển sinh','Chỉ tiêu · số xác nhận · phần còn thiếu'),(4.77,'02','Đào tạo','Cảnh báo · rời trường · ngành và khóa'),(8.84,'03','Lãnh đạo','So sánh · nhận diện điểm cần kiểm tra')]:
    feature(s, x, 4.85, 3.79, n, h, b, h=1.55)

# 03 — current inventory, measured from runtime.
s = slide('Ba trợ lý, ba phạm vi công việc', 'Những gì đã có', 50,
    f'AI Tuyển sinh có {ts["public_documents"]} văn bản công khai và {ts["approved_queries"]} truy vấn dữ liệu định sẵn. '
    f'AI Đào tạo có {dt["public_documents"]} văn bản và {dt["approved_queries"]} truy vấn. Cả hai cơ sở dữ liệu nghiệp vụ hiện đều giả lập. '
    'AI Tổng hợp phục vụ Hiệu trưởng, biết năng lực của các AI con để giao việc và kiểm tra báo cáo. '
    'Con số 106 văn bản, 36 truy vấn trên giao diện Tổng hợp là tổng năng lực của hai AI con, không phải toàn bộ kho được đưa vào AI Tổng hợp.',
    ['backend/app/apps.py:117', 'backend/app/web.py:97', 'CONTENT_FACTS.json'])
for i, (name, audience, stat, body, color) in enumerate([
    ('AI Tuyển sinh','Cán bộ tuyển sinh',f'{ts["public_documents"]} văn bản  /  {ts["approved_queries"]} truy vấn','Tra cứu, soạn nháp\nTheo dõi xác nhận nhập học\nRà soát hồ sơ và số liệu',TEAL),
    ('AI Đào tạo','Phòng Đào tạo · Giảng viên',f'{dt["public_documents"]} văn bản  /  {dt["approved_queries"]} truy vấn','Quy chế và chương trình\nCảnh báo, kết quả học tập\nLớp học phần, lịch thi',AMBER),
    ('AI Tổng hợp','Hiệu trưởng','Giao việc cho 2 AI lĩnh vực','Lập kế hoạch liên lĩnh vực\nĐối chiếu kết quả trả về\nTổng hợp cho lãnh đạo',PURPLE)]):
    x=.7+i*4.07
    rect(s,x,2.46,3.79,3.98,WHITE,LINE)
    pill(s,x+.22,2.72,3.35,audience,SOFT,color,11)
    text(s,x+.23,3.33,3.32,.5,name,25,color,True)
    text(s,x+.23,4.04,3.32,.52,stat,17,INK,True)
    line(s,x+.23,4.73,x+3.55,4.73)
    text(s,x+.23,5.0,3.32,1.07,body,16,MUTED)

# 04 — native editable topology, not a screenshot of a plan.
s = slide('AI Hiệu trưởng thực sự giao việc', 'Cơ chế phối hợp', 70,
    'Trong code hiện tại, AI Tổng hợp có danh sách tài liệu rỗng, không gắn cơ sở dữ liệu trực tiếp. '
    'Nó được giới thiệu năng lực của các AI con, rồi giao yêu cầu cho từng AI. Mỗi AI con tự lập kế hoạch, tra cứu hoặc truy vấn và nghiệm thu riêng. '
    'Cấp trên nhận báo cáo cùng số liệu đã nghiệm thu, có thể giao tác tử Phân tích tính tiếp, rồi đối chiếu và tổng hợp. '
    'Các agent có luồng và ngữ cảnh riêng, có thể dùng cùng họ model; không cần hiểu là ba mô hình được huấn luyện riêng. '
    'Câu hỏi một lĩnh vực có thể chỉ gọi một AI con; bước Phân tích riêng chỉ xuất hiện khi cần.',
    ['backend/app/apps.py:128', 'backend/app/apps.py:89', 'backend/app/workers.py:234', 'backend/app/graph.py:223'], dark=True)
node(s,4.75,2.32,3.83,1.04,'AI Tổng hợp','Nhận câu hỏi và lập kế hoạch','9CBAEE',True)
line(s,6.66,3.36,6.66,3.57,'8BADE6',1.5)
line(s,3.6,3.57,9.72,3.57,'8BADE6',1.5)
line(s,3.6,3.57,3.6,3.86,'8BADE6',1.5,True)
line(s,9.72,3.57,9.72,3.86,'8BADE6',1.5,True)
node(s,1.35,3.87,4.50,1.25,'AI Tuyển sinh','Lập kế hoạch → thực hiện → nghiệm thu','75B79D',True)
node(s,7.47,3.87,4.50,1.25,'AI Đào tạo','Lập kế hoạch → thực hiện → nghiệm thu','DBB074',True)
line(s,3.6,5.12,3.6,5.44,'8BADE6',1.5)
line(s,9.72,5.12,9.72,5.44,'8BADE6',1.5)
line(s,3.6,5.44,9.72,5.44,'8BADE6',1.5)
line(s,6.66,5.44,6.66,5.70,'8BADE6',1.5,True)
rect(s,3.1,5.72,7.13,.62,'304866','577699')
text(s,3.25,5.82,6.83,.35,'Phân tích khi cần → AI Tổng hợp kiểm tra và trả lời',19,WHITE,True,PP_ALIGN.CENTER)
text(s,.79,6.57,11.8,.28,'Cấp Tổng hợp: 0 tài liệu trực tiếp · Không có kết nối CSDL riêng · Nhận kết quả do AI con trả về',12,LIGHT,align=PP_ALIGN.CENTER)

# 05 — recorded execution evidence, exact counts and ordering.
s = slide('Có bằng chứng từ một lượt chạy cụ thể', 'Minh chứng thực thi · 29/09/2026', 50,
    'Bản ghi lúc 02 giờ 51 ngày 29 tháng 9 cho thấy T1 Tuyển sinh và T2 Đào tạo được giao trong cùng một đợt, khoảng giây 40. '
    'Cả hai có kế hoạch riêng. T3 Phân tích khai báo phụ thuộc T1 và T2, chỉ được giao khi hai kết quả đã đạt, khoảng phút 4 giây 13. '
    'Sau đó AI Tổng hợp viết câu trả lời. Tổng 18 lượt gọi dịch vụ AI: Tuyển sinh 8, Đào tạo 4, cấp trên gồm phân tích và tổng hợp 6. '
    'Đây là minh chứng luồng phối hợp, không phải chứng nhận toàn bộ nội dung luôn chính xác.', [trace_source])
rows=[('T1','AI Tuyển sinh','Cùng đợt với T2','8 lượt gọi',TEAL),('T2','AI Đào tạo','Cùng đợt với T1','4 lượt gọi',AMBER),('T3','Tác tử Phân tích','Đợi T1 và T2 đạt','2 lượt gồm thực hiện và nghiệm thu',PURPLE)]
for i,(task,name,dep,calls,color) in enumerate(rows):
    y=2.54+i*1.05
    rect(s,.7,y,8.4,.88,WHITE,LINE)
    circle(s,.9,y+.15,.5,task,color,SOFT,14)
    text(s,1.65,y+.16,3.25,.35,name,19,INK,True)
    text(s,5.12,y+.13,3.7,.28,dep,14,MUTED)
    text(s,5.12,y+.49,3.7,.22,calls,11,color)
metric(s,9.42,2.54,3.20,'18','Lượt gọi AI trong bản ghi')
metric(s,9.42,3.99,3.20,'4 phút 47','Thời gian của riêng lượt này',TEAL)
text(s,.82,5.67,11.73,.24,'18 lượt = 8 Tuyển sinh + 4 Đào tạo + 6 cấp điều phối; trong 6 lượt có 2 lượt của bước Phân tích.',11,MUTED)
rect(s,.72,6.03,11.9,.54,SOFT)
text(s,.91,6.14,11.5,.31,'~00:40  Giao hai AI con     →     ~04:13  Bắt đầu phân tích     →     ~04:47  Trả kết quả',17,BLUE,True)

# 06 — actual current UI; welcome contains no replay switch.
s = slide('Hỏi việc, thấy tiến độ, mở được nguồn', 'Không gian làm việc mới', 35,
    'Đây là giao diện hiện tại. Bên trái chọn trợ lý, ở giữa nhập yêu cầu, bên phải theo dõi các bước. '
    'Kho tài liệu cho tìm theo tên hoặc mã. Nguồn trong câu trả lời mở được văn bản gốc khi có đường dẫn; bảng dữ liệu có thể tải CSV. '
    'Câu nhập tay và gợi ý đều chạy mới, giao diện chính không còn lựa chọn phát lại. Việc đổi trợ lý ở đây chưa phải đăng nhập hay phân quyền người dùng.',
    ['backend/app/pages/index.html','shots/current/welcome-principal.png'])
screen(s,'welcome-principal.png',4.12,2.28,8.5,4.43)
for y,n,head,body in [(2.57,'01','Chọn đúng trợ lý','Tuyển sinh · Đào tạo · Tổng hợp'),(3.98,'02','Nêu một công việc','Tra cứu, phân tích, soạn báo cáo'),(5.39,'03','Kiểm tra căn cứ','Tiến độ, nguồn dẫn và bảng số liệu')]:
    circle(s,.74,y,.43,n,size=13)
    text(s,1.33,y-.02,2.54,.45,head,18,INK,True)
    text(s,1.33,y+.51,2.54,.64,body,14,MUTED)

# 07 — let the presenter actually execute one question; no paid build-time run.
s = slide('Trình diễn một yêu cầu liên lĩnh vực', 'Demo trực tiếp · Dự kiến 5 phút', 300,
    'Chuyển sang trình duyệt tại http://127.0.0.1:8001, chọn AI Tổng hợp. Nhập đúng câu hỏi trên slide. '
    'Đây là chạy mới nên có sử dụng dịch vụ AI, thời gian thay đổi; lượt mẫu đã lưu mất khoảng 4 phút 47 giây. '
    'Trong lúc chờ, giải thích kế hoạch và chỉ rõ việc giao cho Tuyển sinh, Đào tạo, bước phân tích sau khi các đầu vào đạt. '
    'Khi có kết quả, mở tab Nguồn đã dẫn và nhắc số liệu nghiệp vụ là giả lập. Không đọc số liệu như báo cáo thật của Trường. '
    'Nếu dịch vụ lỗi hoặc quá thời gian, chuyển phụ lục B1 và nói rõ đó là ảnh của lượt chạy đã lưu. Nếu trang tự dùng kết quả đã lưu, thông báo đó phải được nói rõ; không giới thiệu là kết quả vừa tính. '
    'Không dùng thao tác Shift+bấm hay tìm nút Bản ghi: các thao tác chọn chế độ này đã bỏ.', [trace_source,'backend/app/pages/index.html'])
rect(s,.7,2.48,11.92,2.12,WHITE,LINE)
pill(s,.97,2.73,2.42,'CHỌN AI TỔNG HỢP',SOFT,PURPLE)
text(s,.97,3.28,11.37,1.06,question,23,INK,True)
for i,(h,b) in enumerate([('1. Xem kế hoạch','Công việc được chia cho AI nào?'),('2. Theo dõi phối hợp','AI con làm gì, kết quả nào đã đạt?'),('3. Mở căn cứ','Nguồn, thời điểm dữ liệu, giới hạn')]):
    x=.73+i*4.07
    text(s,x,5.06,3.7,.44,h,19,BLUE,True)
    text(s,x,5.64,3.7,.65,b,16,MUTED)
text(s,.76,6.53,11.85,.25,'Nếu cần dự phòng: phụ lục B1 là ảnh kết quả đã lưu, không phải một lần chạy mới.',12,MUTED)

# 08/09 — existing examples are honest saved execution evidence, no personal rows.
s = slide('Tuyển sinh: nhìn rõ số liệu cần xử lý', 'Một AI lĩnh vực', 30,
    'Ảnh này lấy từ lượt chạy đã lưu trên bộ dữ liệu giả lập. AI Tuyển sinh truy vấn số xác nhận nhập học và phần chỉ tiêu còn lại theo ngành. '
    'Cán bộ nhận bảng tổng hợp cùng mốc dữ liệu để tiếp tục đối chiếu. Truy vấn được chọn từ danh mục có sẵn, số đếm do cơ sở dữ liệu thực hiện. '
    'Giá trị ở đây là gom một công việc nghiệp vụ thành một luồng kiểm tra được; chưa có số đo độc lập về phần trăm thời gian tiết kiệm.',
    ['backend/app/domains/tuyen_sinh_queries.py','shots/current/screenshots.json'])
screen(s,'answer-admissions.png',.73,2.3,8.28,4.38,'Ảnh từ lượt chạy đã lưu · Số liệu nghiệp vụ giả lập')
text(s,9.37,2.72,3.1,1.10,'Từ câu hỏi\nđến đầu việc',23,TEAL,True)
text(s,9.37,4.04,3.1,1.76,'Số xác nhận theo ngành\nPhần chỉ tiêu còn lại\nMốc dữ liệu rõ ràng\nNguồn để kiểm tra',18,INK)

s = slide('Đào tạo: đặt số liệu trong đúng bối cảnh', 'Một AI lĩnh vực', 30,
    'AI Đào tạo có quy chế, chương trình và bộ dữ liệu học vụ giả lập riêng. Ví dụ này tổng hợp cảnh báo theo ngành và khóa. '
    'Một tỷ lệ chỉ có ý nghĩa khi nêu đúng mẫu số, đúng học kỳ và phạm vi. Vì vậy cán bộ cần kiểm tra cả bảng và nguồn, không chỉ nhìn một con số lớn. '
    'Các truy vấn hiện còn giả định và mốc cố định của bản thử; trước khi nối dữ liệu thật cần hiệu chỉnh với Phòng Đào tạo.',
    ['backend/app/domains/dao_tao_queries.py','shots/current/screenshots.json'])
screen(s,'answer-academic.png',.73,2.3,8.28,4.38,'Ảnh từ lượt chạy đã lưu · Số liệu nghiệp vụ giả lập')
text(s,9.37,2.72,3.1,1.50,'Đúng ngành.\nĐúng khóa.\nĐúng học kỳ.',23,AMBER,True)
text(s,9.37,4.46,3.1,1.6,'Cảnh báo học vụ\nKết quả học tập\nLớp học phần, lịch thi',18,INK)

# 10 — explain the data flow as implemented, not a universal privacy promise.
s = slide('Tách số tổng hợp khỏi danh sách cá nhân', 'Thiết kế luồng dữ liệu', 50,
    'Trong luồng truy vấn hiện tại, model chọn tên truy vấn và tham số; nó không được tự viết SQL để chạy. '
    'Cơ sở dữ liệu mở chỉ đọc. Với truy vấn danh sách cá nhân, phần gửi model là số lượng và thống kê; bảng chi tiết đi tới giao diện. '
    'Điều này chưa phải lời đảm bảo mọi dữ liệu cá nhân không thể ra ngoài: hệ thống vẫn cần đăng nhập, phân quyền nghiệp vụ và kiểm tra từng luồng khi nối dữ liệu thật. '
    'Demo hôm nay chỉ dùng người và hồ sơ giả lập.',
    ['backend/app/tools/data.py:64','backend/app/tools/data.py:113','backend/app/contracts.py:64','backend/app/workers.py:187'])
node(s,.78,3.25,3.2,1.55,'CSDL giả lập','Kết nối chỉ đọc\nTruy vấn được định sẵn',BLUE)
line(s,3.98,3.77,5.0,3.77,BLUE,1.8)
line(s,5.0,3.00,5.0,5.42,BLUE,1.8)
line(s,5.0,3.00,6.1,3.00,TEAL,1.8,True)
line(s,5.0,5.42,6.1,5.42,AMBER,1.8,True)
node(s,6.1,2.49,6.42,1.4,'Số lượng và thống kê → AI','Phân tích, giải thích, soạn báo cáo',TEAL)
node(s,6.1,4.68,6.42,1.4,'Danh sách từng người → Giao diện','Xem bảng và tải CSV; không đưa các hàng này vào prompt',AMBER)
text(s,.81,6.48,11.76,.31,'Cần bổ sung phân quyền người dùng trước khi đưa dữ liệu cá nhân thật vào vận hành.',15,AMBER,True)

# 11 — separate implementation checks from limited model evals.
s = slide('Đã kiểm tra — nhưng chưa phải nghiệm thu độc lập', 'Kết quả và giới hạn', 45,
    '159 kiểm thử phần mềm đã qua trong lần kiểm tra cục bộ ngày 29 tháng 9, bao gồm luồng điều phối, lỗi, truy vấn và khóa cài đặt. '
    'Đó không phải 159 câu hỏi nghiệp vụ với model thật. Các tệp kết quả đã lưu cho thấy 30 trên 30 câu văn bản tuyển sinh, 14 trên 14 câu dữ liệu tuyển sinh giả lập, '
    '10 trên 10 câu Đào tạo và 4 trên 4 câu Tổng hợp đạt bộ chấm nội bộ. Các bộ này đã dùng trong quá trình phát triển; bộ chấm còn thiên về số và mã nguồn xuất hiện. '
    'Vì vậy không suy thành độ chính xác 100% khi dùng thật. Bước tiếp theo phải có câu hỏi mới và người nghiệp vụ chấm độc lập.',
    ['CONTENT_FACTS.json:software_verification','CONTENT_FACTS.json:limited_evaluations'])
rect(s,.73,2.48,3.52,3.75,NAVY)
text(s,1.04,2.83,2.89,1.08,'159',58,WHITE,True)
text(s,1.04,4.11,2.89,.87,'Kiểm thử\nphần mềm đạt',24,WHITE,True)
text(s,1.04,5.40,2.89,.5,'Xác minh cục bộ 29/9/2026',12,LIGHT)
rows=[('Văn bản tuyển sinh','30 / 30'),('Dữ liệu tuyển sinh giả lập','14 / 14'),('Đào tạo giả lập','10 / 10'),('Tổng hợp giả lập','4 / 4')]
for i,(label,value) in enumerate(rows):
    y=2.56+i*.78
    text(s,4.74,y,5.38,.41,label,19,INK)
    text(s,10.65,y,1.56,.41,value,23,BLUE,True,PP_ALIGN.RIGHT)
    line(s,4.76,y+.56,12.27,y+.56)
text(s,4.75,5.97,7.45,.56,'Bộ thử nội bộ đã có đáp án; chưa phải cam kết chất lượng trên dữ liệu thật.',15,AMBER,True)

# 12 — scope honest about settings password vs role auth.
s = slide('Bản thử đã chạy; điều kiện dùng thật còn phải hoàn thiện', 'Phạm vi hiện tại', 45,
    'Nhóm đã có giao diện, ba trợ lý, nguồn dẫn, truy vấn chỉ đọc, theo dõi công việc và mật khẩu cho trang cài đặt. '
    'Mật khẩu cài đặt chỉ bảo vệ cấu hình kết nối; nó không thay thế đăng nhập và phân quyền cho từng phòng ban. '
    'Hiện chưa nối cơ sở dữ liệu nghiệp vụ thật, chưa phục hồi công việc đang chạy sau khi tiến trình sập và chưa có cơ chế hủy hoàn chỉnh. '
    'Vì vậy đề nghị bước thử nghiệm nhỏ có kiểm soát, chưa đề nghị mở dùng toàn trường.', ['CONTENT_FACTS.json:current_boundaries'])
rect(s,.75,2.47,5.76,3.94,WHITE,LINE)
rect(s,6.80,2.47,5.76,3.94,WHITE,LINE)
pill(s,1.02,2.75,2.63,'ĐÃ CÓ TRONG CODE',PALE_GREEN,TEAL)
pill(s,7.07,2.75,3.85,'TRƯỚC KHI THỬ DỮ LIỆU THẬT',PALE_AMBER,AMBER)
text(s,1.04,3.46,5.13,2.37,'Ba trợ lý và cơ chế giao việc\nTra nguồn, truy vấn, bảng kết quả\nNhật ký và chi phí lượt chạy\nMật khẩu cho trang cài đặt',20,INK)
text(s,7.09,3.46,5.13,2.37,'Đăng nhập, phân quyền nghiệp vụ\nDữ liệu được đơn vị duyệt\nPhục hồi, hủy và vận hành bền\nNghiệm thu bằng câu hỏi mới',20,INK)
text(s,1.04,6.0,10.95,.27,'Mật khẩu cài đặt ≠ phân quyền dữ liệu theo phòng ban.',14,AMBER,True)

# 13 — measured records, not current vendor pricing.
s = slide('Chi phí phụ thuộc độ phức tạp của yêu cầu', 'Ước tính trong hồ sơ chạy', 40,
    'Chi phí đang nói là ước tính từ mức dùng và bảng giá cấu hình tại thời điểm chạy, không phải báo giá dịch vụ hiện tại. '
    'Trong bộ 14 câu dữ liệu tuyển sinh đã lưu, trung bình khoảng 0,026 đô la, xấp xỉ 700 đồng; thời gian trung vị khoảng 43 giây. '
    'Riêng câu Tổng hợp liên lĩnh vực đã trình bày, ước tính 0,265 đô la, khoảng 6.900 đồng và mất 4 phút 47 giây. '
    'Không lấy hai con số này để cam kết chi phí mỗi câu hay suy ra kinh phí toàn trường. Khi pilot cần đo trên công việc thật, theo dõi cả chất lượng và chi phí.',
    ['CONTENT_FACTS.json:limited_evaluations','CONTENT_FACTS.json:principal_example','backend/app/config.py:PRICE_DATE'])
for x, label, amount, latency, detail, color in [
    (.75,'14 CÂU DỮ LIỆU TUYỂN SINH','~700 đ / câu','Trung vị ~43 giây','Trung bình ước tính $0,026\nBộ thử nội bộ trên dữ liệu giả lập',TEAL),
    (6.80,'1 CÂU TỔNG HỢP LIÊN LĨNH VỰC','~6.900 đ','4 phút 47 giây','Ước tính $0,265 · 18 lượt gọi AI\nRiêng bản ghi minh chứng đã trình bày',BLUE)]:
    rect(s,x,2.5,5.76,3.51,WHITE,LINE)
    text(s,x+.29,2.81,5.17,.31,label,12,color,True)
    text(s,x+.27,3.50,5.2,.79,amount,37,color,True)
    text(s,x+.29,4.49,5.17,.49,latency,20,INK,True)
    text(s,x+.29,5.16,5.17,.66,detail,15,MUTED)
text(s,.81,6.29,11.78,.45,'Quy đổi minh họa: 26.180 đ/USD theo cấu hình demo. Số liệu là ước tính đã ghi ngày 29/9, không phải báo giá hiện tại.',12,MUTED)

# 14 — sequence of proposed pilot gates, not an unsupported delivery promise.
s = slide('Bước tiếp theo: thử nhỏ, đo thật, rồi mới mở rộng', 'Lộ trình đề xuất', 40,
    'Nhóm đề nghị đi theo ba bước có điều kiện hoàn thành. Trước hết chọn một số công việc và đầu mối nghiệp vụ, xác định dữ liệu được phép dùng. '
    'Sau đó bổ sung phân quyền, phục hồi và nối bản sao chỉ đọc trong phạm vi được duyệt. Cuối cùng dùng câu hỏi mới, so với đáp án cán bộ, đo thời gian và chi phí. '
    'Chỉ mở rộng khi đạt các tiêu chí đã thống nhất. Đây là trình tự đề xuất; thời gian cụ thể phụ thuộc dữ liệu, hạ tầng và người nghiệm thu.', ['PLAN.md','DATA_PLAN.md','CONTENT_FACTS.json:current_boundaries'])
for i,(num,head,body,gate) in enumerate([
    ('01','Chốt công việc','Chọn tình huống ưu tiên\nCử người phụ trách nguồn\nXác định quyền sử dụng','Đầu vào có người xác nhận'),
    ('02','Hoàn thiện nền tảng','Đăng nhập và phân quyền\nBản sao dữ liệu chỉ đọc\nPhục hồi và vận hành','Qua kiểm tra quyền, lỗi'),
    ('03','Nghiệm thu thử','Câu hỏi mới, chấm độc lập\nĐo chất lượng, thời gian, phí\nGhi nhận phản hồi thực tế','Đạt tiêu chí mới mở rộng')]):
    x=.75+i*4.04
    rect(s,x,2.57,3.70,3.95,WHITE,LINE)
    text(s,x+.24,2.82,3.16,.80,num,40,'9BB2D0',True)
    text(s,x+.24,3.8,3.16,.52,head,23,INK,True)
    text(s,x+.24,4.54,3.16,1.25,body,17,MUTED)
    pill(s,x+.22,5.93,3.26,gate,SOFT,BLUE,10)

# 15 — leave the meeting with a concrete request.
s = slide('Để bắt đầu, nhóm cần ba sự phối hợp', 'Đề nghị Nhà trường', 40,
    'Nhóm xin đề nghị ba việc. Một, cử đầu mối Tuyển sinh và Đào tạo để chọn công việc đáng làm và kiểm tra kết quả. '
    'Hai, thống nhất một tập dữ liệu được phép dùng và cách tạo bản sao chỉ đọc; chuẩn bị tài khoản dịch vụ AI do Trường quản lý trước khi dùng dữ liệu thật. '
    'Ba, cung cấp một bộ câu hỏi mới cùng đáp án hoặc tiêu chí để nghiệm thu độc lập. Nhóm sẽ báo lại những gì làm được, những gì chưa đạt, thời gian và chi phí đo thực tế. '
    'Em xin cảm ơn quý thầy cô. Công việc nào của đơn vị nên được chọn để thử đầu tiên?', ['DATA_PLAN.md','PLAN.md'], dark=True)
for y,n,h,b in [(2.53,'01','Hai đầu mối nghiệp vụ','Tuyển sinh và Đào tạo cùng chọn việc, xác nhận nguồn.'),
                 (3.76,'02','Một phạm vi dữ liệu được duyệt','Bản sao chỉ đọc, quyền sử dụng và tài khoản dịch vụ rõ ràng.'),
                 (4.99,'03','Một bộ câu hỏi nghiệm thu mới','Có đáp án hoặc tiêu chí do cán bộ chuyên môn xác nhận.')]:
    circle(s,.79,y,.57,n,'D7E5F8','304866',18)
    text(s,1.65,y-.03,10.46,.53,h,25,WHITE,True)
    text(s,1.65,y+.61,10.46,.42,b,18,LIGHT)
text(s,.82,6.46,11.75,.35,'Bắt đầu từ công việc cụ thể. Đo được giá trị trước khi mở rộng.',21,'A8C8F2',True)

# Backup B1 — explain what the screenshot really represents.
s = slide('B1 · Kết quả Tổng hợp đã lưu để dự phòng', 'Ảnh minh chứng', 0,
    'Chỉ dùng khi không chạy được demo hoặc cần xem lại kết quả. Đây là ảnh giao diện hiện tại đang hiển thị kết quả của lượt chạy đã lưu lúc 02:51 ngày 29/9/2026. '
    'Số liệu nội bộ là giả lập; không trình bày đây là báo cáo thật hay kết quả vừa chạy. Bản ghi nêu ngành Kế toán và có các giới hạn về mẫu số, phạm vi dữ liệu. '
    'Chỉ vào cả phần bảng lẫn phần lưu ý, không chỉ đọc kết luận.', [trace_source,'shots/current/answer-principal.png'], backup=True)
screen(s,'answer-principal.png',.8,2.26,11.74,4.42,'Ảnh kết quả từ hồ sơ 29/9/2026 · Văn bản công khai + số liệu nghiệp vụ giả lập')

s = slide('B2 · AI Tổng hợp nhận gì và không nhận gì?', 'Giải đáp về kiến trúc', 0,
    'AI Tổng hợp lúc lập kế hoạch chỉ biết yêu cầu và mô tả năng lực AI con, gồm loại số liệu có thể truy vấn. '
    'Khi AI con làm xong, nó nhận báo cáo, nhận định có mã nguồn và các kết quả có cấu trúc đã nghiệm thu. '
    'Nó không được truyền toàn bộ nội dung tài liệu hay cơ sở dữ liệu của hai AI con. Tác tử phân tích tính trên các đầu vào đã được trả về. '
    'Đây là cách nối dữ liệu trong luồng ứng dụng, chưa phải ranh giới cách ly hạ tầng hay phân quyền người dùng hoàn chỉnh. '
    'Cùng một model có thể đóng nhiều vai với các ngữ cảnh và lời gọi riêng.',
    ['backend/app/apps.py:82','backend/app/apps.py:128','backend/app/workers.py:234','backend/app/graph.py:223'], backup=True)
items=[('Lập kế hoạch','Yêu cầu + mô tả năng lực các AI con'),('AI con trả về','Báo cáo + nguồn dẫn + số liệu đã nghiệm thu'),('Phân tích','Kết quả tác vụ trước; phép tính do code thực hiện'),('Kho trực tiếp của Tổng hợp','0 tài liệu · không gắn DataStore')]
for i,(a,b) in enumerate(items):
    y=2.5+i*.91
    rect(s,.8,y,11.73,.75,WHITE,LINE)
    text(s,1.04,y+.16,3.71,.42,a,18,BLUE,True)
    text(s,5.05,y+.16,7.1,.42,b,18,INK)
text(s,.87,6.44,11.64,.34,'Bằng chứng có trong mã nguồn, kế hoạch tác vụ, sự kiện AI con và sổ ghi lời gọi.',15,MUTED)

s = slide('B3 · Căn cứ nội dung và cách kiểm chứng', 'Phụ lục cho người xem kỹ thuật', 0,
    'Các số liệu và giới hạn trong bài được tổng hợp trong CONTENT_FACTS.json, có đường dẫn nguồn cho từng nhận định. '
    'Ảnh chụp và metadata nằm trong shots/current. Bản ghi minh chứng có kế hoạch riêng của hai AI con, thứ tự giao việc và 18 lời gọi AI. '
    'Kết quả bộ thử nghiệp vụ là tệp đã lưu; lượt làm lại slide này không gọi model thật hoặc chạy lại bộ đánh giá có phí. '
    'File PowerPoint chứa lời nói và nguồn trong phần Notes. Kịch bản Markdown đi kèm ghi rõ cách demo và xử lý tình huống dự phòng.',
    ['CONTENT_FACTS.json','shots/current/screenshots.json','KICH_BAN_TRINH_BAY.md'], backup=True)
for y,h,b in [(2.54,'Cơ chế giao việc','backend/app/apps.py · workers.py · graph.py'),
               (3.40,'Dữ liệu và truy vấn','backend/app/domains/ · tools/data.py'),
               (4.26,'Kết quả đã lưu','eval-20260929-0149 · du-lieu-20260929-0138 · tong-hop-20260929-0156'),
               (5.12,'Hồ sơ bài trình bày','CONTENT_FACTS.json · shots/current/screenshots.json · Notes của từng slide')]:
    text(s,.84,y,3.4,.45,h,18,BLUE,True)
    text(s,4.1,y,8.4,.54,b,16,INK)
    line(s,.84,y+.64,12.43,y+.64)
text(s,.85,6.26,11.5,.4,'Cập nhật theo code và dữ liệu đang có ngày 29/9/2026. Không chứa mật khẩu hoặc khóa kết nối.',13,MUTED)


def write_script():
    total = sum(i['seconds'] for i in INDEX)
    assert total == 900, total
    lines = ['# Credere — kịch bản trình bày cho Ban Giám hiệu', '',
             'Cập nhật 29/09/2026. 15 slide chính, 3 phụ lục dự phòng. Tổng thời lượng dự kiến: **15 phút**, gồm khoảng 5 phút cho một lượt demo.', '',
             '**Thông điệp:** AI Tổng hợp giao việc thật cho AI Tuyển sinh và AI Đào tạo; kết quả phục vụ người có chuyên môn kiểm tra và quyết định. Văn bản là nguồn công khai, số liệu nội bộ trong demo là giả lập.', '',
             '## Trước buổi trình bày', '',
             '- Mở `Credere_Trinh_bay_BGH.pptx` bằng PowerPoint và bật Presenter View để xem Notes.',
             '- Mở web tại http://127.0.0.1:8001. Nếu cần khởi động: chạy `C:\\Users\\laqdv\\.credere\\venv\\Scripts\\python.exe -m app.web --port 8001` trong thư mục `credere/backend`.',
             '- Chọn AI Tổng hợp và kiểm tra kho tài liệu, nhãn dữ liệu giả lập. Không thay dữ liệu hoặc chạy lại bộ đánh giá ngay trước buổi trình bày.',
             '- Trang cài đặt yêu cầu mật khẩu; chỉ người quản trị mở khi cần. Không trình chiếu mật khẩu hoặc khóa kết nối.',
             '- Câu nhập tay và câu gợi ý mặc định chạy mới. Không có nút Bản ghi/Trực tiếp hoặc phím Shift+bấm để đổi chế độ.',
             '- Nếu dịch vụ lỗi hoặc hết thời gian, dùng slide B1, nói rõ là ảnh kết quả đã lưu. Không giới thiệu ảnh cũ như kết quả vừa tính.', '',
             '## Thời lượng', '', '| Slide | Nội dung | Dự kiến |', '| --- | --- | --- |']
    for i in INDEX:
        if not i['backup']:
            sec=i['seconds']; lines.append(f'| {i["slide"]} | {i["title"]} | {sec//60}:{sec%60:02d} |')
    lines += ['', '## Lời nói theo slide', '']
    for i in INDEX:
        lines += [f'### {i["slide"]}. {i["title"]}', '', i['notes'], '']
        if i['sources']:
            lines += ['Căn cứ: ' + '; '.join(f'`{src}`' for src in i['sources']), '']
    lines += ['## Hỏi đáp thường gặp', '',
              '| Câu hỏi | Trả lời gợi ý |', '| --- | --- |',
              '| AI Hiệu trưởng có đọc toàn bộ kho của hai AI kia không? | Không trong luồng hiện tại: `sources=[]`, `data=None`. Nó giao việc, nhận báo cáo và kết quả đã nghiệm thu; có thể dùng tác tử Phân tích để tính tiếp. |',
              '| Hai AI con có thực sự chạy riêng? | Có lời gọi graph và client riêng, kế hoạch/sự kiện riêng. Có thể dùng cùng model ở các ngữ cảnh khác nhau. |',
              '| Câu nào cũng phải gọi đủ hai AI con? | Không. Tùy công việc: một lĩnh vực có thể chỉ gọi một, nhiệm vụ liên lĩnh vực có thể song song hoặc phụ thuộc nhau. |',
              '| Vì sao không có câu trả lời ngay? | Một yêu cầu gồm lập kế hoạch, truy vấn, nghiệm thu, có thể sửa và tổng hợp. Bản ghi liên lĩnh vực mất 4 phút 47 giây; đó không phải cam kết thời gian. |',
              '| Có thể dùng dữ liệu thật ngay chưa? | Chưa. Cần quyền sử dụng dữ liệu, đăng nhập/phân quyền nghiệp vụ, hiệu chỉnh truy vấn và nghiệm thu độc lập. Mật khẩu cài đặt chỉ bảo vệ cấu hình. |',
              '| Các kết quả 30/30, 14/14 có nghĩa AI luôn đúng? | Không. Đó là bộ thử nội bộ đã dùng trong phát triển, bộ chấm còn giới hạn. Cán bộ kiểm tra nguồn và cần câu hỏi mới để nghiệm thu. |',
              '| Chi phí khoảng 700 đồng là cố định? | Không. Là trung bình ước tính của bộ thử số liệu cụ thể; câu phức tạp có nhiều lượt gọi sẽ khác. |',
              '| AI có xét tuyển hoặc xử lý học vụ thay cán bộ? | Không. AI hỗ trợ tra cứu, thống kê và soạn nháp; quyết định do người có trách nhiệm. |', '',
              '## Các tệp bàn giao', '',
              '- `Credere_Trinh_bay_BGH.pptx`: bản chính, nội dung chỉnh sửa được, Notes có lời nói và căn cứ.',
              '- `Credere_Trinh_bay_BGH.pdf`: bản xem và chia sẻ.',
              '- `CONTENT_FACTS.json`: kiểm chứng số liệu/nhận định từ code và hồ sơ chạy.',
              '- `shots/current/`: ảnh giao diện hiện tại; metadata ghi nguồn và thời điểm của kết quả đã lưu.',
              '- `archive/before-redesign-20260929/`: bản trình bày cũ được giữ lại.']
    (HERE/'KICH_BAN_TRINH_BAY.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    (HERE/'deck_numbers.json').write_text(json.dumps({'updated_on':'2026-09-29','runtime':FACTS['runtime'],
        'software_tests':159,'limited_evaluations':FACTS['limited_evaluations'],'principal_example':example},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def audit():
    errors=[]
    for n,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            if sh.left < -100 or sh.top < -100 or sh.left+sh.width>prs.slide_width+100 or sh.top+sh.height>prs.slide_height+100:
                errors.append(f'Slide {n}: shape outside page: {sh.name}')
    notes=[s.notes_slide.notes_text_frame.text for s in prs.slides]
    if not all(len(n)>80 for n in notes): errors.append('Missing substantial speaker notes')
    report={'slides':len(prs.slides),'main_slides':15,'backup_slides':3,'speaker_seconds':sum(i['seconds'] for i in INDEX),
            'shape_bounds_errors':errors,'all_slides_have_notes':all(len(n)>80 for n in notes),'text_boxes':TEXT_BOXES,'index':INDEX}
    (HERE/'deck_review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    if errors: raise ValueError('\n'.join(errors))


for slide_ in prs.slides:
    for shape in slide_.shapes:
        for effect_ref in shape._element.xpath('.//a:effectRef'):
            effect_ref.set('idx', '0')  # LibreOffice otherwise inherits the template's drop shadow.

write_script()
audit()
prs.save(OUT)
shutil.copy2(OUT, HERE/'AI_Tuyen_sinh_trinh_bay.pptx')  # keep the existing user's entry point up to date
print(f'Built {OUT.name}: {len(prs.slides)} slides (15 main + 3 backups), 15-minute script, editable diagrams and speaker notes.')
