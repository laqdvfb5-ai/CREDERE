"""Render every slide and ensure slide text survives PDF export."""
from pathlib import Path
import json
import re
import unicodedata
import pypdfium2 as pdfium
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE=Path(__file__).resolve().parent
OUT=HERE/'render'/'15p'
NAME='Credere_15_phut_day_du'
doc=pdfium.PdfDocument(str(OUT/f'{NAME}.pdf'))
deck=Presentation(HERE/f'{NAME}.pptx')
assert len(doc)==len(deck.slides)==23
font=ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf',16)
norm=lambda s:re.sub(r'\s+','',unicodedata.normalize('NFKC',s)).replace('\u00ad','')
thumbs=[];warnings=[];bounds=[]
for i,page in enumerate(doc):
    pic=page.render(scale=1.5).to_pil().convert('RGB')
    pic.save(OUT/f'slide-{i+1:02}.png')
    actual=norm(page.get_textpage().get_text_range())
    for sh in deck.slides[i].shapes:
        if sh.has_text_frame and norm(sh.text) and norm(sh.text) not in actual:warnings.append({'slide':i+1,'text':sh.text})
    pic.thumbnail((480,270))
    cell=Image.new('RGB',(500,306),'#e8edf4');cell.paste(pic,((500-pic.width)//2,8))
    ImageDraw.Draw(cell).text((12,281),f'{i+1:02}',fill='#344968',font=font);thumbs.append(cell)
for start in range(0,len(thumbs),9):
    sheet=Image.new('RGB',(1500,918),'#e8edf4')
    for j,pic in enumerate(thumbs[start:start+9]):sheet.paste(pic,((j%3)*500,(j//3)*306))
    sheet.save(OUT/f'overview-{start//9+1}.png')
(OUT/'render_review.json').write_text(json.dumps({'pages':len(doc),'text_warnings':warnings},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'pages':len(doc),'text_warnings':len(warnings)}))
