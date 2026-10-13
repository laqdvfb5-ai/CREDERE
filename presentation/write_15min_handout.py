"""Create an editable Word handout from the same reviewed slide content."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parent
content=json.loads((ROOT/'review_15p.json').read_text(encoding='utf-8'))
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
style=doc.styles['Normal']
style.font.name='Calibri';style.font.size=Pt(11)
style.paragraph_format.space_after=Pt(7)
for name in ('Title','Heading 1','Heading 2','Heading 3'):
    doc.styles[name].font.name='Calibri'
    doc.styles[name].font.color.rgb=RGBColor.from_string('245BBA')
doc.add_heading('CREDERE',0)
doc.add_heading('Nội dung chi tiết bài thuyết trình 15 phút',1)
doc.add_paragraph('Theo yeucauslide.txt · Cập nhật 29/09/2026')
doc.add_paragraph('15 slide chính, tổng 900 giây; 8 phụ lục chỉ dùng khi hỏi đáp. Tổng quan mô tả toàn bộ công việc dự án sẽ làm: sáu lĩnh vực, nhóm người dùng, nguồn/đầu ra, phối hợp liên phòng và sản phẩm bàn giao. Lời nói tách khỏi phần giải thích để giữ thời lượng. Bản demo là mục minh chứng riêng ở slide 6, dành tối đa 1 phút 20 giây.')
doc.add_paragraph('Demo: https://proxyllm.store. Văn bản là nguồn công khai; hai CSDL nghiệp vụ vẫn giả lập. Dự toán local dùng model lớn và giá GPU đã kiểm tra: bốn card Max-Q 96GB riêng gần 2 tỷ; tổng máy/hạ tầng chờ báo giá. Chi vận hành có kịch bản riêng, không đồng nghĩa đã đo khả năng phục vụ hoặc chất lượng tương đương API.')
doc.add_heading('Phân bổ thời gian',1)
table=doc.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
for c,t in zip(table.rows[0].cells,['Slide','Nội dung','Thời lượng']):c.text=t
for i,s in enumerate(content['slides'][:15],1):
    cells=table.add_row().cells
    for c,t in zip(cells,[str(i),s['title'],f'{s["seconds"]//60}:{s["seconds"]%60:02d}']):c.text=t
for i,s in enumerate(content['slides'],1):
    doc.add_page_break()
    doc.add_heading(f'Slide {i}. {s["title"]}',1)
    doc.add_paragraph('Phụ lục hỏi đáp' if s['backup'] else f'Thời lượng dự kiến: {s["seconds"]} giây')
    doc.add_heading('Thông điệp và nội dung trên slide',2)
    doc.add_paragraph(s['takeaway'])
    if s['headers']:
        table=doc.add_table(rows=1,cols=len(s['headers']));table.style='Light Shading Accent 1'
        for c,t in zip(table.rows[0].cells,s['headers']):c.text=t
        for row in s['rows']:
            for c,t in zip(table.add_row().cells,row):c.text=t
    else:
        for title,body in s['items']:doc.add_paragraph(f'{title}: {body}',style='List Bullet')
    if s['foot']:doc.add_paragraph(s['foot'])
    if s['image']:doc.add_paragraph(f'Ảnh: shots/current/{s["image"]}')
    doc.add_heading('Lời thuyết trình',2)
    doc.add_paragraph(s['notes'])
    if s.get('explanation'):
        doc.add_heading('Giải thích bổ sung để hỏi đáp',2)
        doc.add_paragraph('Không đọc toàn bộ phần này trong 15 phút.')
        doc.add_paragraph(s['explanation'])
    doc.add_heading('Nguồn và căn cứ',3)
    for source in s['sources']:doc.add_paragraph(source)
doc.core_properties.title='Credere — Nội dung chi tiết bài thuyết trình 15 phút'
doc.core_properties.author='Nhóm đề tài Credere'
doc.save(ROOT/'NOI_DUNG_CHI_TIET_15P.docx')
print('Word handout created: 15 timed slides + 8 optional appendices.')
