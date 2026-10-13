"""Build the editable Credere presentation and the matching speaker script.

Run with ~/.credere/venv/Scripts/python.exe presentation/build_deck.py.
Evidence: CONTENT_FACTS.json; screenshots: shots/current/screenshots.json.
No model calls or network access are needed to build this deck.
"""
import json
import shutil
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE / 'shots' / 'current'
FACTS = json.loads((HERE / 'CONTENT_FACTS.json').read_text(encoding='utf-8'))
OUT = HERE / 'Credere_Chi_tiet_du_an.pptx'
FONT = 'Segoe UI'
NAVY, INK, MUTED = '142238', '23364F', '687A91'
BLUE, TEAL, AMBER, PURPLE = '245BBA', '40836F', 'A76B27', '7755A2'
PAPER, WHITE, LINE, SOFT = 'F6F7F9', 'FFFFFF', 'DFE5ED', 'EBF1F9'
LIGHT, PALE_GREEN, PALE_AMBER = 'AFBFDA', 'ECF5F0', 'FBF3E6'
DATE = '29.09.2026'
W, H = 13.333333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Credere — Trợ lý AI phối hợp cho Trường Đại học Nghệ An'
prs.core_properties.subject = 'Bản demo hiện tại, cơ chế giao việc và đề xuất thử nghiệm có kiểm soát'
prs.core_properties.author = 'Nhóm đề tài Credere'
prs.core_properties.keywords = 'Credere, NAU, tuyển sinh, đào tạo, đa tác tử'
INDEX = []
TEXT_BOXES = []


def rgb(value):
    return RGBColor.from_string(value)


def rect(s, x, y, w, h, fill=WHITE, line=None, radius=.12):
    obj = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    if line:
        obj.line.color.rgb = rgb(line)
        obj.line.width = Pt(.8)
    else:
        obj.line.fill.background()
    if radius:
        obj.adjustments[0] = radius
    obj.shadow.inherit = False
    return obj


def text(s, x, y, w, h, content, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=.025):
    obj = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    lines = content.split('\n') if isinstance(content, str) else content
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.1
        p.space_after = Pt(5 if len(lines) > 1 else 0)
        r = p.add_run()
        r.text = str(line)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    TEXT_BOXES.append({'slide': len(prs.slides), 'text': '\n'.join(lines), 'font': size,
                       'x': x, 'y': y, 'w': w, 'h': h})
    return obj


def line(s, x1, y1, x2, y2, color=LINE, width=1.4, arrow=False):
    obj = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.line.color.rgb, obj.line.width = rgb(color), Pt(width)
    obj.shadow.inherit = False
    if arrow:
        end = etree.SubElement(obj.line._get_or_add_ln(), qn('a:tailEnd'))
        end.set('type', 'triangle')
    return obj


def pill(s, x, y, w, label, fill=SOFT, color=BLUE, size=12):
    rect(s, x, y, w, .34, fill=fill, radius=.2)
    text(s, x+.09, y+.025, w-.18, .25, label, size, color, bold=True)


def circle(s, x, y, d, label, color=BLUE, fill=SOFT, size=18):
    obj = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    obj.fill.solid(); obj.fill.fore_color.rgb = rgb(fill); obj.line.fill.background()
    obj.shadow.inherit = False
    text(s, x, y, d, d, label, size, color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def brand(s, x=.6, y=.22, color=NAVY):
    text(s, x, y, 2.2, .38, 'credere.', 20, color, bold=True)


def slide(title, section, seconds, notes, sources=(), dark=False, backup=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY if dark else PAPER)
    fg, muted = (WHITE, LIGHT) if dark else (INK, MUTED)
    brand(s, color=WHITE if dark else NAVY)
    text(s, 7.8, .29, 4.9, .23, f'TRƯỜNG ĐẠI HỌC NGHỆ AN  /  {DATE}', 10, muted, align=PP_ALIGN.RIGHT)
    line(s, .62, .82, 12.71, .82, '34465F' if dark else LINE, .7)
    text(s, .64, 1.03, 12, .28, section.upper(), 11, '98B5E7' if dark else BLUE, bold=True)
    if title:
        text(s, .62, 1.43, 12.1, .73, title, 31, fg, bold=True)
    num = len(prs.slides)
    text(s, .64, 7.02, 10.9, .24, 'Credere · Hệ thống AI đa tác tử cho nhà trường' + (' · PHỤ LỤC DỰ PHÒNG' if backup else ''), 10, muted)
    text(s, 12.05, 6.98, .65, .3, f'{num:02}', 12, muted, align=PP_ALIGN.RIGHT)
    note = notes + ('\n\nCăn cứ (dành cho người trình bày):\n' + '\n'.join(str(p) for p in sources) if sources else '')
    s.notes_slide.notes_text_frame.text = note
    INDEX.append({'slide':num,'title':title or 'Credere — trợ lý AI phối hợp','section':section,'seconds':seconds,
                  'notes':notes,'sources':list(sources),'backup':backup})
    return s


def metric(s, x, y, w, number, label, color=BLUE, fill=WHITE):
    rect(s, x, y, w, 1.20, fill, LINE)
    text(s, x+.2, y+.13, w-.4, .55, str(number), 28, color, bold=True)
    text(s, x+.2, y+.77, w-.4, .32, label, 12, MUTED)


def feature(s, x, y, w, number, title, body, color=BLUE, fill=WHITE, h=1.3):
    rect(s, x, y, w, h, fill, LINE)
    circle(s, x+.18, y+.19, .42, number, color, SOFT if color==BLUE else PALE_GREEN, 13)
    text(s, x+.78, y+.15, w-.98, .45, title, 18, INK, bold=True)
    text(s, x+.78, y+.68, w-.98, h-.78, body, 14, MUTED)


def screen(s, filename, x, y, w, h, caption=None):
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(f'Missing actual screenshot: {path}')
    im = Image.open(path)
    iw, ih = im.size
    scale = min(w/iw, h/ih)
    pw, ph = iw*scale, ih*scale
    rect(s, x+(w-pw)/2-.035, y+(h-ph)/2-.035, pw+.07, ph+.07, WHITE, LINE, radius=.03)
    s.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))
    if caption:
        text(s, x, y+h+.10, w, .45, caption, 10, MUTED)


def node(s, x, y, w, h, title, body, color=BLUE, dark=False):
    fill, fg, secondary = ('223750', WHITE, LIGHT) if dark else (WHITE, INK, MUTED)
    rect(s, x, y, w, h, fill, '4B6380' if dark else LINE)
    rect(s, x+.02, y+.13, .045, h-.26, color, radius=False)
    text(s, x+.18, y+.12, w-.36, .42, title, 19, fg, bold=True)
    text(s, x+.18, y+.64, w-.36, h-.72, body, 14, secondary)


