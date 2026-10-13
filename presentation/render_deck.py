"""Render the exported PDF for visual review; verify page count and retained text."""
import json
import re
from pathlib import Path
import unicodedata

import pypdfium2 as pdfium
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
OUT = HERE/'render'/'current'
PDF = OUT/'Credere_Trinh_bay_BGH.pdf'
doc = pdfium.PdfDocument(str(PDF))
deck = Presentation(HERE/'Credere_Trinh_bay_BGH.pptx')
assert len(doc) == len(deck.slides) == 18
thumbs=[]
missing=[]
font=ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf',16)

def norm(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', value)).replace('\u00ad','')

for i,page in enumerate(doc):
    bitmap=page.render(scale=1.5)
    pic=bitmap.to_pil().convert('RGB')
    pic.save(OUT/f'slide-{i+1:02}.png')
    txt=page.get_textpage()
    actual=norm(txt.get_text_range())
    for shape in deck.slides[i].shapes:
        if shape.has_text_frame:
            expected=norm(shape.text)
            if expected and expected not in actual:
                missing.append({'slide':i+1,'text':shape.text})
    pic.thumbnail((480,270))
    cell=Image.new('RGB',(500,306),'#e8edf4')
    cell.paste(pic,((500-pic.width)//2,8))
    ImageDraw.Draw(cell).text((12,281),f'{i+1:02}',fill='#344968',font=font)
    thumbs.append(cell)

for start in (0,9):
    sheet=Image.new('RGB',(1500,918),'#e8edf4')
    for j,thumb in enumerate(thumbs[start:start+9]): sheet.paste(thumb,((j%3)*500,(j//3)*306))
    sheet.save(OUT/f'overview-{start//9+1}.png')
report={'pdf_pages':len(doc),'powerpoint_slides':len(deck.slides),'text_not_found_in_pdf':missing}
(OUT/'render_review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'pages':len(doc),'text_warnings':len(missing)}))
